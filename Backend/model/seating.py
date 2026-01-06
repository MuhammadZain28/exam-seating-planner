from openpyxl import Workbook
from openpyxl.styles import Border, Side, Alignment, PatternFill
from io import BytesIO
from .rooms import Rooms
from .exam import Exams
from ..structures.Queue import PriorityQueue
import json

class Seating:
    _instance = None
    def __new__(cls):
        if cls._instance is None:
            cls._instance = super(Seating, cls).__new__(cls)
            cls._instance.arrangements = {}
        return cls._instance

    def arrange_exam_hall(self, rooms, exams):
        self.radixSort(rooms)
        halls = {}
        count = 0
        while not exams.isEmpty() and rooms:
            room = rooms.pop(0)
            halls[room["name"]] = {}
            courses = {}
            section = {'A': 0, 'B': 0, 'C': 0, 'D': 0}
            groups = {0: [], 1: [], 2: [], 3: []}

            for r in range(room["rows"]):
                for c in range(room["columns"]):
                    g = (r % 2) * 2 + (c % 2)
                    groups[g].append((r, c))


            hall = [[None for _ in range(room["columns"])] for _ in range(room["rows"])]

            i = 0
            remaining_seats = []
            new_exams = []
            while i < 4 and not exams.isEmpty():
                exam_info = exams.pop()
                if not exam_info:
                    break
                key = exam_info['session']
                courses[key] = {}

                for student in exam_info["students"]:
                    if not groups[i]:
                        remaining_seats.append(student)
                        continue
                    r, c = groups[i].pop()
                    hall[r][c] = (student["reg"])
                    section[student.get("section", "A")] += 1

                    count += 1
                if remaining_seats and len(rooms):
                    new_exams.append({"course": exam_info["course"], "date": exam_info["date"], "duration": exam_info["duration"], "session": exam_info["session"], "students": remaining_seats})
                    remaining_seats = []
                courses[key] = section.copy()
                section = {'A': 0, 'B': 0, 'C': 0, 'D': 0}
                i = (i + 1) % 4
            if new_exams:
                for ne in new_exams:
                    exams.push(ne, len(exam_info["students"]))

            halls[room["name"]] = {"layout": hall[::-1], "courses": courses}
        print("Total Students Seated:", count)
        return halls

    def radixSort(self, arr):
        max1 = max(arr["rows"] * arr["columns"] for arr in arr)
        exp = 1
        while max1 // exp > 0:
            self.countingSort(arr, exp)
            exp *= 10
        arr.reverse()
    def countingSort(self, arr, exp):
        n = len(arr)
        output = [0] * n
        count = [0] * 10

        for i in range(n):
            index = (arr[i]["rows"] * arr[i]["columns"] // exp) % 10
            count[index] += 1

        for i in range(1, 10):
            count[i] += count[i - 1]

        for i in range(n - 1, -1, -1):
            index = (arr[i]["rows"] * arr[i]["columns"] // exp) % 10
            output[count[index] - 1] = arr[i]
            count[index] -= 1

        for i in range(n):
            arr[i] = output[i]

    def shedule_exams(self, date="2025-12-21", time="09:00"):
        roomInstance = Rooms()
        examInstance = Exams()
        rooms = roomInstance.to_list()
        exams = PriorityQueue()
        for exam in examInstance.get():
            if exam["date"] == date and exam["time"] == time:
                exams.push(exam, len(exam["students"]))
            elif exam["date"] > date:
                break
        arrangement = self.arrange_exam_hall(rooms, exams)
        self.save_arrangement(arrangement, date, time)
        return arrangement

    def save_arrangement(self, arrangement, date, time):
        self.arrangements.setdefault(date, {})[time] = arrangement
        with open("Exam_arrangements.json", "w") as f:
            json.dump(self.arrangements, f, indent=1)

    def load_arrangement(self):
        try:
            filename = f"Exam_arrangements.json"
            with open(filename, "r") as f:
                self.arrangements = json.load(f)
        except FileNotFoundError:
            return None

    def load_exist(self, date, time):
        if date in self.arrangements and time in self.arrangements[date]:
            return self.arrangements[date][time]
        return None

    def get_arrangements(self):
        data = []
        for date, times in self.arrangements.items():
            for time, arrangements in times.items():
                data.append({
                    "date": date,
                    "time": time
                })
        return data

    def delete_arrangement(self, date, time):
        if date in self.arrangements and time in self.arrangements[date]:
            del self.arrangements[date][time]
            if not self.arrangements[date]:
                del self.arrangements[date]
            with open("Exam_arrangements.json", "w") as f:
                json.dump(self.arrangements, f, indent=1)
            return True
        return False

    def export_xlsx(self, date, time):
        data = self.load_exist(date, time)

        wb = Workbook()
        wb.remove(wb.active)

        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        center_align = Alignment(horizontal="center", vertical="center")

        batch_fills = {
            "2021": PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid"),
            "2022": PatternFill(start_color="C6E0B4", end_color="C6E0B4", fill_type="solid"),
            "2023": PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),
            "2024": PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid"),
        }

        for hall, details in data.items():
            ws = wb.create_sheet(title=hall)

            layout = details["layout"]
            courses = details["courses"]  # batch -> section -> count

            # ---------------- Seating Layout ----------------
            for r, row in enumerate(layout, start=1):
                for c, value in enumerate(row, start=1):
                    cell = ws.cell(row=r, column=c, value=value)
                    cell.border = thin_border
                    cell.alignment = center_align

                    if isinstance(value, str):
                        batch = value.split("-")[0]
                        if batch in batch_fills:
                            cell.fill = batch_fills[batch]

            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = 18

            # ---------------- Summary Table ----------------
            start_row = len(layout) + 2

            title_cell = ws.cell(row=start_row, column=1, value="Courses Summary")
            title_cell.alignment = center_align
            title_cell.border = thin_border

            headers = ["Batch", "Section", "Count"]
            for col, header in enumerate(headers, start=1):
                cell = ws.cell(row=start_row + 1, column=col, value=header)
                cell.alignment = center_align
                cell.border = thin_border

            row_no = start_row + 2

            for batch, sections in courses.items():
                for section, count in sections.items():
                    b_cell = ws.cell(row=row_no, column=1, value=batch)
                    s_cell = ws.cell(row=row_no, column=2, value=section)
                    c_cell = ws.cell(row=row_no, column=3, value=count)

                    for cell in (b_cell, s_cell, c_cell):
                        cell.alignment = center_align
                        cell.border = thin_border

                    if batch in batch_fills:
                        b_cell.fill = batch_fills[batch]
                        s_cell.fill = batch_fills[batch]
                        c_cell.fill = batch_fills[batch]

                    row_no += 1

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        return output

